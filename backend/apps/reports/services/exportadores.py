"""Exportacion de los reportes a Excel y PDF.

Ambos exportadores consumen la misma estructura que produce `generador.py`
(titulo, columnas, filas, totales), de modo que agregar un reporte nuevo no
obliga a tocar este modulo.
"""

import io
from datetime import date, datetime, time
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AZUL = "1F4E79"
GRIS = "F2F2F2"


def _texto(valor):
    """Representacion legible de un valor para la celda del reporte."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%Y-%m-%d")
    if isinstance(valor, time):
        return valor.strftime("%H:%M")
    if isinstance(valor, Decimal):
        return f"{valor:.2f}"
    return str(valor)


def exportar_excel(datos, fecha_inicio, fecha_fin, subtitulo=""):
    """Genera el archivo Excel del reporte y lo devuelve en memoria."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = datos["titulo"][:31]

    columnas = datos["columnas"]
    total_columnas = len(columnas)

    # --- Encabezado ---------------------------------------------------
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columnas)
    celda = hoja.cell(row=1, column=1, value=datos["titulo"])
    celda.font = Font(size=14, bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)
    celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.row_dimensions[1].height = 24

    hoja.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columnas)
    periodo = f"Periodo: {fecha_inicio} al {fecha_fin}"
    if subtitulo:
        periodo = f"{periodo}  |  {subtitulo}"
    celda = hoja.cell(row=2, column=1, value=periodo)
    celda.alignment = Alignment(horizontal="center")
    celda.font = Font(size=10, italic=True)

    # --- Cabecera de la tabla -----------------------------------------
    fila_cabecera = 4
    borde = Border(*[Side(style="thin", color="BFBFBF")] * 4)
    for indice, (_, etiqueta) in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_cabecera, column=indice, value=etiqueta)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde

    # --- Datos ---------------------------------------------------------
    for numero, fila in enumerate(datos["filas"], start=fila_cabecera + 1):
        for indice, (clave, _) in enumerate(columnas, start=1):
            valor = fila.get(clave)
            celda = hoja.cell(row=numero, column=indice)
            # Los numeros se escriben como numeros para que Excel pueda sumarlos.
            if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
                celda.value = float(valor)
                celda.number_format = "#,##0.00" if isinstance(valor, (float, Decimal)) else "#,##0"
            else:
                celda.value = _texto(valor)
            celda.border = borde
            if numero % 2 == 0:
                celda.fill = PatternFill("solid", fgColor=GRIS)

    # --- Totales -------------------------------------------------------
    if datos.get("totales"):
        fila_totales = fila_cabecera + len(datos["filas"]) + 1
        hoja.cell(row=fila_totales, column=1, value="TOTALES").font = Font(bold=True)
        for indice, (clave, _) in enumerate(columnas, start=1):
            if clave in datos["totales"]:
                celda = hoja.cell(row=fila_totales, column=indice)
                celda.value = float(datos["totales"][clave])
                celda.font = Font(bold=True)
                celda.number_format = "#,##0.00"

    if datos.get("nota"):
        fila_nota = fila_cabecera + len(datos["filas"]) + 3
        hoja.merge_cells(
            start_row=fila_nota, start_column=1, end_row=fila_nota, end_column=total_columnas
        )
        celda = hoja.cell(row=fila_nota, column=1, value=f"Nota: {datos['nota']}")
        celda.font = Font(size=9, italic=True)
        celda.alignment = Alignment(wrap_text=True)

    _ajustar_anchos(hoja, columnas, datos["filas"])
    hoja.freeze_panes = hoja.cell(row=fila_cabecera + 1, column=1)

    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return buffer


def _ajustar_anchos(hoja, columnas, filas):
    """Ajusta el ancho de cada columna al contenido mas largo."""
    for indice, (clave, etiqueta) in enumerate(columnas, start=1):
        ancho = len(etiqueta)
        for fila in filas[:200]:  # una muestra basta para estimar el ancho
            ancho = max(ancho, len(_texto(fila.get(clave))))
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(ancho + 3, 10), 40)


def exportar_pdf(datos, fecha_inicio, fecha_fin, subtitulo=""):
    """Genera el PDF del reporte y lo devuelve en memoria."""
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1 * cm,
        rightMargin=1 * cm,
        topMargin=1 * cm,
        bottomMargin=1 * cm,
        title=datos["titulo"],
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(datos["titulo"], estilos["Title"]),
        Paragraph(
            f"Periodo: {fecha_inicio} al {fecha_fin}" + (f" | {subtitulo}" if subtitulo else ""),
            estilos["Normal"],
        ),
        Spacer(1, 0.5 * cm),
    ]

    columnas = datos["columnas"]
    tabla_datos = [[etiqueta for _, etiqueta in columnas]]
    for fila in datos["filas"]:
        tabla_datos.append([_texto(fila.get(clave)) for clave, _ in columnas])

    if datos.get("totales"):
        fila_totales = ["TOTALES"] + [""] * (len(columnas) - 1)
        for indice, (clave, _) in enumerate(columnas):
            if clave in datos["totales"]:
                fila_totales[indice] = _texto(datos["totales"][clave])
        tabla_datos.append(fila_totales)

    tabla = Table(tabla_datos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{GRIS}")]),
    ]
    if datos.get("totales"):
        estilo.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    tabla.setStyle(TableStyle(estilo))

    elementos.append(tabla)

    if datos.get("nota"):
        elementos.append(Spacer(1, 0.5 * cm))
        elementos.append(Paragraph(f"<i>Nota: {datos['nota']}</i>", estilos["Normal"]))

    documento.build(elementos)
    buffer.seek(0)
    return buffer
